
import os
import shutil
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference


# ============================================================
# CONFIGURAÇÕES
# ============================================================

pasta_entrada = "entrada"
pasta_saida = "saida"
pasta_processados = "processados"

arquivo_historico = os.path.join(
    pasta_saida,
    "dados_historicos.csv"
)

arquivo_relatorio = os.path.join(
    pasta_saida,
    "relatorio_vendas.xlsx"
)

arquivo_log = os.path.join(
    pasta_saida,
    "log_execucoes.csv"
)

arquivo_controle = os.path.join(
    pasta_saida,
    "controle_arquivos.csv"
)

colunas_obrigatorias = [
    "Data",
    "Cliente",
    "Produto",
    "Quantidade",
    "Valor"
]


# ============================================================
# PREPARAÇÃO DAS PASTAS
# ============================================================

os.makedirs(pasta_entrada, exist_ok=True)
os.makedirs(pasta_saida, exist_ok=True)
os.makedirs(pasta_processados, exist_ok=True)


# ============================================================
# FORMATAÇÃO
# ============================================================

def ajustar_larguras(ws):

    larguras = {
        "A": 18,
        "B": 20,
        "C": 22,
        "D": 18,
        "E": 18,
        "F": 18,
        "G": 20,
        "H": 18,
        "I": 18
    }

    for coluna, largura in larguras.items():
        ws.column_dimensions[coluna].width = largura


def formatar_planilha(ws):

    for cell in ws[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    for row in ws.iter_rows(
        min_row=2
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="center"
            )

    ajustar_larguras(ws)


# ============================================================
# LEITURA E VALIDAÇÃO
# ============================================================

def ler_e_validar_arquivo(
    caminho,
    nome_arquivo,
    todos_erros
):

    erros_arquivo = []

    try:

        dados = pd.read_excel(
            caminho,
            engine="openpyxl"
        )

    except Exception as erro:

        erros_arquivo.append({
            "Arquivo": nome_arquivo,
            "Tipo de Erro": "Erro ao ler arquivo",
            "Detalhes": str(erro)
        })

        todos_erros.extend(
            erros_arquivo
        )

        return None

    # --------------------------------------------------------
    # Remove espaços extras dos nomes das colunas
    # --------------------------------------------------------

    dados.columns = (
        dados.columns
        .astype(str)
        .str.strip()
    )

    print(
        "COLUNAS LIDAS PELO PYTHON:"
    )

    print(
        dados.columns.tolist()
    )

    # --------------------------------------------------------
    # Verifica colunas obrigatórias
    # --------------------------------------------------------

    colunas_faltantes = [
        coluna
        for coluna in colunas_obrigatorias
        if coluna not in dados.columns
    ]

    if colunas_faltantes:

        erros_arquivo.append({
            "Arquivo": nome_arquivo,
            "Tipo de Erro": "Colunas obrigatórias ausentes",
            "Detalhes": (
                "Colunas obrigatórias ausentes: "
                + ", ".join(colunas_faltantes)
            )
        })

        todos_erros.extend(
            erros_arquivo
        )

        return None

    dados = dados[
        colunas_obrigatorias
    ].copy()

    # --------------------------------------------------------
    # Arquivo vazio
    # --------------------------------------------------------

    if dados.empty:

        erros_arquivo.append({
            "Arquivo": nome_arquivo,
            "Tipo de Erro": "Arquivo vazio",
            "Detalhes": (
                "O arquivo não possui registros."
            )
        })

        todos_erros.extend(
            erros_arquivo
        )

        return None

    # --------------------------------------------------------
    # Limpeza de textos
    # --------------------------------------------------------

    dados["Cliente"] = (
        dados["Cliente"]
        .astype(str)
        .str.strip()
    )

    dados["Produto"] = (
        dados["Produto"]
        .astype(str)
        .str.strip()
    )

    # --------------------------------------------------------
    # Campos vazios
    # --------------------------------------------------------

    for coluna in colunas_obrigatorias:

        quantidade_vazios = int(
            dados[coluna].isna().sum()
        )

        if quantidade_vazios > 0:

            erros_arquivo.append({
                "Arquivo": nome_arquivo,
                "Tipo de Erro": "Campos vazios",
                "Detalhes": (
                    f"{coluna}: "
                    f"{quantidade_vazios} "
                    "registro(s) vazio(s)"
                )
            })

    # --------------------------------------------------------
    # Data
    # --------------------------------------------------------

    data_original = dados["Data"].copy()

    dados["Data"] = pd.to_datetime(
        dados["Data"],
        errors="coerce",
        dayfirst=True
    )

    datas_invalidas = (
        dados["Data"].isna()
        & data_original.notna()
    )

    quantidade_datas_invalidas = int(
        datas_invalidas.sum()
    )

    if quantidade_datas_invalidas > 0:

        erros_arquivo.append({
            "Arquivo": nome_arquivo,
            "Tipo de Erro": "Data inválida",
            "Detalhes": (
                f"{quantidade_datas_invalidas} "
                "registro(s) com data inválida"
            )
        })

    # --------------------------------------------------------
    # Quantidade
    # --------------------------------------------------------

    quantidade_original = (
        dados["Quantidade"].copy()
    )

    dados["Quantidade"] = pd.to_numeric(
        dados["Quantidade"],
        errors="coerce"
    )

    quantidades_invalidas = (
        dados["Quantidade"].isna()
        & quantidade_original.notna()
    )

    quantidade_erros = int(
        quantidades_invalidas.sum()
    )

    if quantidade_erros > 0:

        erros_arquivo.append({
            "Arquivo": nome_arquivo,
            "Tipo de Erro": "Quantidade inválida",
            "Detalhes": (
                f"{quantidade_erros} "
                "registro(s) com quantidade inválida"
            )
        })

    # --------------------------------------------------------
    # Valor
    # --------------------------------------------------------

    valor_original = (
        dados["Valor"].copy()
    )

    dados["Valor"] = pd.to_numeric(
        dados["Valor"],
        errors="coerce"
    )

    valores_invalidos = (
        dados["Valor"].isna()
        & valor_original.notna()
    )

    quantidade_valores_invalidos = int(
        valores_invalidos.sum()
    )

    if quantidade_valores_invalidos > 0:

        erros_arquivo.append({
            "Arquivo": nome_arquivo,
            "Tipo de Erro": "Valor inválido",
            "Detalhes": (
                f"{quantidade_valores_invalidos} "
                "registro(s) com valor inválido"
            )
        })

    # --------------------------------------------------------
    # Remove registros inválidos
    # --------------------------------------------------------

    antes = len(dados)

    dados = dados.dropna(
        subset=colunas_obrigatorias
    ).copy()

    linhas_removidas = (
        antes - len(dados)
    )

    if linhas_removidas > 0:

        erros_arquivo.append({
            "Arquivo": nome_arquivo,
            "Tipo de Erro": "Registros inválidos ignorados",
            "Detalhes": (
                f"{linhas_removidas} "
                "registro(s) inválido(s) "
                "foram ignorados"
            )
        })

    # --------------------------------------------------------
    # Cliente e produto vazios
    # --------------------------------------------------------

    dados = dados[
        (dados["Cliente"] != "")
        & (dados["Cliente"].str.lower() != "nan")
        & (dados["Produto"] != "")
        & (dados["Produto"].str.lower() != "nan")
    ].copy()

    # --------------------------------------------------------
    # Nenhum registro válido
    # --------------------------------------------------------

    if dados.empty:

        erros_arquivo.append({
            "Arquivo": nome_arquivo,
            "Tipo de Erro": "Nenhum registro válido",
            "Detalhes": (
                "O arquivo não possui registros "
                "válidos para processamento."
            )
        })

        todos_erros.extend(
            erros_arquivo
        )

        return None

    # --------------------------------------------------------
    # Duplicidades dentro do arquivo
    # --------------------------------------------------------

    duplicados = dados.duplicated(
        subset=colunas_obrigatorias,
        keep="first"
    )

    quantidade_duplicados = int(
        duplicados.sum()
    )

    if quantidade_duplicados > 0:

        erros_arquivo.append({
            "Arquivo": nome_arquivo,
            "Tipo de Erro": "Registros duplicados",
            "Detalhes": (
                f"{quantidade_duplicados} "
                "registro(s) duplicado(s)"
            )
        })

        dados = dados[
            ~duplicados
        ].copy()

    todos_erros.extend(
        erros_arquivo
    )

    return dados


# ============================================================
# RESUMO POR CLIENTE
# ============================================================

def gerar_resumo_por_cliente(dados):

    resumo = (
        dados
        .groupby("Cliente")
        .agg({
            "Total": "sum",
            "Quantidade": "sum"
        })
        .reset_index()
    )

    resumo = resumo.rename(
        columns={
            "Total": "Total Comprado (R$)",
            "Quantidade": "Quantidade Total"
        }
    )

    produtos_por_cliente = []

    for cliente in dados["Cliente"].unique():

        dados_cliente = dados[
            dados["Cliente"] == cliente
        ]

        quantidades = (
            dados_cliente
            .groupby("Produto")["Quantidade"]
            .sum()
            .sort_values(
                ascending=False
            )
        )

        produto_mais_comprado = (
            quantidades.index[0]
        )

        quantidade_produto = (
            quantidades.iloc[0]
        )

        produtos_por_cliente.append({
            "Cliente": cliente,
            "Produto Mais Comprado":
                produto_mais_comprado,
            "Quantidade":
                quantidade_produto
        })

    produtos = pd.DataFrame(
        produtos_por_cliente
    )

    resumo = resumo.merge(
        produtos,
        on="Cliente",
        how="left"
    )

    return resumo[
        [
            "Cliente",
            "Total Comprado (R$)",
            "Produto Mais Comprado",
            "Quantidade"
        ]
    ]


# ============================================================
# CONTROLE DE ARQUIVOS
# ============================================================

def registrar_controle_arquivo(
    nome_arquivo,
    data_processamento,
    status,
    linhas_lidas,
    linhas_novas,
    duplicidades,
    quantidade_erros,
    destino
):

    registro = pd.DataFrame([
        {
            "Arquivo": nome_arquivo,
            "Data/Hora": data_processamento,
            "Status": status,
            "Linhas lidas": linhas_lidas,
            "Linhas novas": linhas_novas,
            "Duplicidades ignoradas": duplicidades,
            "Erros": quantidade_erros,
            "Destino": destino
        }
    ])

    if os.path.exists(
        arquivo_controle
    ):

        registro.to_csv(
            arquivo_controle,
            mode="a",
            header=False,
            index=False,
            encoding="utf-8-sig"
        )

    else:

        registro.to_csv(
            arquivo_controle,
            index=False,
            encoding="utf-8-sig"
        )


# ============================================================
# GERAÇÃO DO RELATÓRIO
# ============================================================

def gerar_relatorio(
    dados,
    todos_erros
):

    dados = dados.copy()

    dados["Total"] = (
        dados["Quantidade"]
        * dados["Valor"]
    )

    # --------------------------------------------------------
    # Vendas por mês
    # --------------------------------------------------------

    dados["Mês"] = (
        dados["Data"]
        .dt.strftime("%m/%Y")
    )

    vendas_mes = (
        dados
        .groupby("Mês")["Total"]
        .sum()
        .reset_index()
    )

    vendas_mes = vendas_mes.rename(
        columns={
            "Total": "Total de Vendas"
        }
    )

    # --------------------------------------------------------
    # Vendas por produto
    # --------------------------------------------------------

    vendas_produto = (
        dados
        .groupby("Produto")["Total"]
        .sum()
        .reset_index()
    )

    vendas_produto = vendas_produto.rename(
        columns={
            "Total": "Total de Vendas"
        }
    )

    # --------------------------------------------------------
    # Vendas por cliente
    # --------------------------------------------------------

    vendas_cliente = (
        dados
        .groupby("Cliente")["Total"]
        .sum()
        .reset_index()
    )

    vendas_cliente = vendas_cliente.rename(
        columns={
            "Total": "Total de Vendas"
        }
    )

    # --------------------------------------------------------
    # Resumo por cliente
    # --------------------------------------------------------

    resumo_cliente = (
        gerar_resumo_por_cliente(dados)
    )

    # --------------------------------------------------------
    # Controle de arquivos
    # --------------------------------------------------------

    if os.path.exists(
        arquivo_controle
    ):

        controle_arquivos = pd.read_csv(
            arquivo_controle,
            encoding="utf-8-sig"
        )

    else:

        controle_arquivos = pd.DataFrame(
            columns=[
                "Arquivo",
                "Data/Hora",
                "Status",
                "Linhas lidas",
                "Linhas novas",
                "Duplicidades ignoradas",
                "Erros",
                "Destino"
            ]
        )

    # --------------------------------------------------------
    # Relatório de erros
    # --------------------------------------------------------

    if todos_erros:

        relatorio_erros = pd.DataFrame(
            todos_erros
        )

    else:

        relatorio_erros = pd.DataFrame(
            columns=[
                "Arquivo",
                "Tipo de Erro",
                "Detalhes"
            ]
        )

    # --------------------------------------------------------
    # Salva Excel
    # --------------------------------------------------------

    with pd.ExcelWriter(
        arquivo_relatorio,
        engine="openpyxl"
    ) as writer:

        dados.to_excel(
            writer,
            sheet_name="Dados Consolidados",
            index=False
        )

        vendas_mes.to_excel(
            writer,
            sheet_name="Vendas por Mês",
            index=False
        )

        vendas_produto.to_excel(
            writer,
            sheet_name="Vendas por Produto",
            index=False
        )

        vendas_cliente.to_excel(
            writer,
            sheet_name="Vendas por Cliente",
            index=False
        )

        resumo_cliente.to_excel(
            writer,
            sheet_name="Resumo por Cliente",
            index=False
        )

        controle_arquivos.to_excel(
            writer,
            sheet_name="Controle de Arquivos",
            index=False
        )

        relatorio_erros.to_excel(
            writer,
            sheet_name="Relatório de Erros",
            index=False
        )

    # --------------------------------------------------------
    # Formatação
    # --------------------------------------------------------

    wb = load_workbook(
        arquivo_relatorio
    )

    for ws in wb.worksheets:

        formatar_planilha(ws)

    # --------------------------------------------------------
    # Dashboard
    # --------------------------------------------------------

    if "Dashboard" in wb.sheetnames:

        del wb["Dashboard"]

    dashboard = wb.create_sheet(
        "Dashboard",
        0
    )

    dashboard["A1"] = (
        "DASHBOARD DE VENDAS"
    )

    dashboard["A1"].font = Font(
        bold=True,
        size=20
    )

    total_vendas = dados["Total"].sum()

    numero_vendas = len(dados)

    clientes = dados[
        "Cliente"
    ].nunique()

    produtos = dados[
        "Produto"
    ].nunique()

    dashboard["A3"] = (
        "TOTAL DE VENDAS"
    )

    dashboard["B3"] = total_vendas

    dashboard["A4"] = (
        "NÚMERO DE VENDAS"
    )

    dashboard["B4"] = numero_vendas

    dashboard["A5"] = (
        "CLIENTES"
    )

    dashboard["B5"] = clientes

    dashboard["A6"] = (
        "PRODUTOS"
    )

    dashboard["B6"] = produtos

    for linha in range(3, 7):

        dashboard[
            f"A{linha}"
        ].font = Font(
            bold=True
        )

        dashboard[
            f"B{linha}"
        ].font = Font(
            bold=True,
            size=14
        )

    dashboard["B3"].number_format = (
        "#,##0.00"
    )

    dashboard.column_dimensions[
        "A"
    ].width = 25

    dashboard.column_dimensions[
        "B"
    ].width = 20

    # --------------------------------------------------------
    # Gráfico por mês
    # --------------------------------------------------------

    if len(vendas_mes) > 0:

        chart_mes = BarChart()

        chart_mes.title = (
            "Vendas por Mês"
        )

        chart_mes.y_axis.title = (
            "Valor"
        )

        chart_mes.x_axis.title = (
            "Mês"
        )

        dados_grafico = Reference(
            wb["Vendas por Mês"],
            min_col=2,
            min_row=1,
            max_row=len(vendas_mes) + 1
        )

        categorias = Reference(
            wb["Vendas por Mês"],
            min_col=1,
            min_row=2,
            max_row=len(vendas_mes) + 1
        )

        chart_mes.add_data(
            dados_grafico,
            titles_from_data=True
        )

        chart_mes.set_categories(
            categorias
        )

        chart_mes.height = 8
        chart_mes.width = 14

        dashboard.add_chart(
            chart_mes,
            "D3"
        )

    # --------------------------------------------------------
    # Gráfico por produto
    # --------------------------------------------------------

    if len(vendas_produto) > 0:

        chart_produto = BarChart()

        chart_produto.title = (
            "Vendas por Produto"
        )

        chart_produto.y_axis.title = (
            "Valor"
        )

        chart_produto.x_axis.title = (
            "Produto"
        )

        dados_grafico = Reference(
            wb["Vendas por Produto"],
            min_col=2,
            min_row=1,
            max_row=len(vendas_produto) + 1
        )

        categorias = Reference(
            wb["Vendas por Produto"],
            min_col=1,
            min_row=2,
            max_row=len(vendas_produto) + 1
        )

        chart_produto.add_data(
            dados_grafico,
            titles_from_data=True
        )

        chart_produto.set_categories(
            categorias
        )

        chart_produto.height = 8
        chart_produto.width = 14

        dashboard.add_chart(
            chart_produto,
            "D20"
        )

    # --------------------------------------------------------
    # Gráfico por cliente
    # --------------------------------------------------------

    if len(vendas_cliente) > 0:

        chart_cliente = BarChart()

        chart_cliente.title = (
            "Vendas por Cliente"
        )

        chart_cliente.y_axis.title = (
            "Valor"
        )

        chart_cliente.x_axis.title = (
            "Cliente"
        )

        dados_grafico = Reference(
            wb["Vendas por Cliente"],
            min_col=2,
            min_row=1,
            max_row=len(vendas_cliente) + 1
        )

        categorias = Reference(
            wb["Vendas por Cliente"],
            min_col=1,
            min_row=2,
            max_row=len(vendas_cliente) + 1
        )

        chart_cliente.add_data(
            dados_grafico,
            titles_from_data=True
        )

        chart_cliente.set_categories(
            categorias
        )

        chart_cliente.height = 8
        chart_cliente.width = 14

        dashboard.add_chart(
            chart_cliente,
            "T3"
        )

    wb.save(
        arquivo_relatorio
    )


# ============================================================
# LOG GERAL
# ============================================================

def registrar_log(
    arquivos_processados,
    arquivos_rejeitados,
    linhas_novas,
    linhas_consolidadas,
    total_vendas
):

    registro = pd.DataFrame([
        {
            "Data/Hora": datetime.now().strftime(
                "%d/%m/%Y %H:%M:%S"
            ),
            "Arquivos processados":
                arquivos_processados,
            "Arquivos rejeitados":
                arquivos_rejeitados,
            "Linhas novas":
                linhas_novas,
            "Linhas consolidadas":
                linhas_consolidadas,
            "Total das vendas":
                total_vendas
        }
    ])

    if os.path.exists(
        arquivo_log
    ):

        registro.to_csv(
            arquivo_log,
            mode="a",
            header=False,
            index=False,
            encoding="utf-8-sig"
        )

    else:

        registro.to_csv(
            arquivo_log,
            index=False,
            encoding="utf-8-sig"
        )


# ============================================================
# MOVER PARA PROCESSADOS
# ============================================================

def mover_para_processados(
    arquivo
):

    origem = os.path.join(
        pasta_entrada,
        arquivo
    )

    destino = os.path.join(
        pasta_processados,
        arquivo
    )

    contador = 1

    nome, extensao = os.path.splitext(
        arquivo
    )

    while os.path.exists(destino):

        novo_nome = (
            f"{nome}_{contador}{extensao}"
        )

        destino = os.path.join(
            pasta_processados,
            novo_nome
        )

        contador += 1

    try:

        shutil.move(
            origem,
            destino
        )

        return destino

    except Exception as erro:

        print(
            f"Erro ao mover {arquivo}: "
            f"{erro}"
        )

        return ""


# ============================================================
# PROGRAMA PRINCIPAL
# ============================================================

print()

print(
    "INICIANDO AUTOMAÇÃO"
)

print(
    "=" * 50
)

arquivos = sorted([
    arquivo
    for arquivo in os.listdir(
        pasta_entrada
    )
    if arquivo.lower().endswith(".xlsx")
    and not arquivo.startswith("~$")
])


# ============================================================
# NENHUM ARQUIVO
# ============================================================

if not arquivos:

    print()

    print(
        "NENHUM ARQUIVO NOVO"
    )

    print(
        "A pasta entrada não possui "
        "arquivos para processar."
    )

    print()

    raise SystemExit


lista_dados = []

arquivos_validos = []

arquivos_rejeitados = []

todos_erros = []

linhas_lidas_total = 0

data_processamento = (
    datetime.now().strftime(
        "%d/%m/%Y %H:%M:%S"
    )
)


# ============================================================
# PROCESSAMENTO
# ============================================================

for arquivo in arquivos:

    print()

    print(
        f"Processando: {arquivo}"
    )

    caminho = os.path.join(
        pasta_entrada,
        arquivo
    )

    quantidade_erros_antes = len(
        todos_erros
    )

    dados = ler_e_validar_arquivo(
        caminho,
        arquivo,
        todos_erros
    )

    erros_deste_arquivo = (
        len(todos_erros)
        - quantidade_erros_antes
    )

    # --------------------------------------------------------
    # Arquivo rejeitado
    # --------------------------------------------------------

    if dados is None:

        arquivos_rejeitados.append(
            arquivo
        )

        print()

        print(
            f"ARQUIVO REJEITADO: {arquivo}"
        )

        for erro in todos_erros[
            quantidade_erros_antes:
        ]:

            print(
                f"- {erro['Tipo de Erro']}: "
                f"{erro['Detalhes']}"
            )

        registrar_controle_arquivo(
            nome_arquivo=arquivo,
            data_processamento=
                data_processamento,
            status="REJEITADO",
            linhas_lidas=0,
            linhas_novas=0,
            duplicidades=0,
            quantidade_erros=
                erros_deste_arquivo,
            destino=""
        )

        continue

    # --------------------------------------------------------
    # Arquivo válido
    # --------------------------------------------------------

    linhas_lidas = len(
        dados
    )

    linhas_lidas_total += (
        linhas_lidas
    )

    lista_dados.append(
        dados
    )

    arquivos_validos.append(
        arquivo
    )


# ============================================================
# NENHUM ARQUIVO VÁLIDO
# ============================================================

if not lista_dados:

    print()

    print(
        "=" * 50
    )

    print(
        "NENHUM ARQUIVO VÁLIDO"
    )

    print(
        "=" * 50
    )

    print(
        f"Arquivos encontrados: "
        f"{len(arquivos)}"
    )

    print(
        f"Arquivos rejeitados: "
        f"{len(arquivos_rejeitados)}"
    )

    if todos_erros:

        caminho_erros = os.path.join(
            pasta_saida,
            "relatorio_erros.xlsx"
        )

        pd.DataFrame(
            todos_erros
        ).to_excel(
            caminho_erros,
            index=False,
            engine="openpyxl"
        )

    print()

    raise SystemExit


# ============================================================
# NOVOS DADOS
# ============================================================

dados_novos = pd.concat(
    lista_dados,
    ignore_index=True
)


# ============================================================
# HISTÓRICO
# ============================================================

if os.path.exists(
    arquivo_historico
):

    try:

        historico = pd.read_csv(
            arquivo_historico,
            encoding="utf-8-sig"
        )

        historico["Data"] = (
            pd.to_datetime(
                historico["Data"],
                errors="coerce"
            )
        )

    except Exception:

        historico = pd.DataFrame(
            columns=colunas_obrigatorias
        )

else:

    historico = pd.DataFrame(
        columns=colunas_obrigatorias
    )


# ============================================================
# TIPOS DO HISTÓRICO
# ============================================================

if not historico.empty:

    historico["Cliente"] = (
        historico["Cliente"]
        .astype(str)
        .str.strip()
    )

    historico["Produto"] = (
        historico["Produto"]
        .astype(str)
        .str.strip()
    )

    historico["Quantidade"] = (
        pd.to_numeric(
            historico["Quantidade"],
            errors="coerce"
        )
    )

    historico["Valor"] = (
        pd.to_numeric(
            historico["Valor"],
            errors="coerce"
        )
    )


# ============================================================
# CONSOLIDA
# ============================================================

linhas_historico_antes = len(
    historico
)

dados_completos = pd.concat(
    [
        historico,
        dados_novos
    ],
    ignore_index=True
)


# ============================================================
# DUPLICIDADES
# ============================================================

antes_duplicidades = len(
    dados_completos
)

dados_completos = (
    dados_completos
    .drop_duplicates(
        subset=colunas_obrigatorias,
        keep="first"
    )
    .reset_index(drop=True)
)

duplicidades_ignoradas = (
    antes_duplicidades
    - len(dados_completos)
)


# ============================================================
# LINHAS NOVAS
# ============================================================

linhas_consolidadas = len(
    dados_completos
)

linhas_novas = (
    linhas_consolidadas
    - linhas_historico_antes
)

if linhas_novas < 0:

    linhas_novas = 0


# ============================================================
# REGISTRA DUPLICIDADES
# ============================================================

if duplicidades_ignoradas > 0:

    todos_erros.append({
        "Arquivo": "Processamento geral",
        "Tipo de Erro": "Registros duplicados",
        "Detalhes": (
            f"{duplicidades_ignoradas} "
            "registro(s) duplicado(s) "
            "ignorado(s)"
        )
    })


# ============================================================
# DISTRIBUI LINHAS NOVAS / DUPLICIDADES POR ARQUIVO
# ============================================================

linhas_restantes = linhas_novas

for arquivo in arquivos_validos:

    caminho = os.path.join(
        pasta_entrada,
        arquivo
    )

    try:

        dados_arquivo = pd.read_excel(
            caminho,
            engine="openpyxl"
        )

        dados_arquivo.columns = (
            dados_arquivo.columns
            .astype(str)
            .str.strip()
        )

        if all(
            coluna in dados_arquivo.columns
            for coluna in colunas_obrigatorias
        ):

            quantidade_linhas = len(
                dados_arquivo
            )

        else:

            quantidade_linhas = 0

    except Exception:

        quantidade_linhas = 0

    linhas_novas_arquivo = min(
        quantidade_linhas,
        linhas_restantes
    )

    linhas_restantes -= (
        linhas_novas_arquivo
    )

    erros_arquivo = sum(
        1
        for erro in todos_erros
        if erro["Arquivo"] == arquivo
    )

    duplicidades_arquivo = 0

    for erro in todos_erros:

        if (
            erro["Arquivo"] == arquivo
            and erro["Tipo de Erro"]
            == "Registros duplicados"
        ):

            texto = erro["Detalhes"]

            try:

                numero = int(
                    texto.split()[0]
                )

                duplicidades_arquivo += (
                    numero
                )

            except Exception:

                pass

    # --------------------------------------------------------
    # Move arquivo
    # --------------------------------------------------------

    destino = mover_para_processados(
        arquivo
    )

    registrar_controle_arquivo(
        nome_arquivo=arquivo,
        data_processamento=
            data_processamento,
        status="PROCESSADO",
        linhas_lidas=quantidade_linhas,
        linhas_novas=linhas_novas_arquivo,
        duplicidades=
            duplicidades_arquivo,
        quantidade_erros=erros_arquivo,
        destino=destino
    )


# ============================================================
# SALVA HISTÓRICO
# ============================================================

dados_historico_salvar = (
    dados_completos[
        colunas_obrigatorias
    ].copy()
)

dados_historico_salvar.to_csv(
    arquivo_historico,
    index=False,
    encoding="utf-8-sig"
)


# ============================================================
# GERA RELATÓRIO
# ============================================================

gerar_relatorio(
    dados_historico_salvar,
    todos_erros
)


# ============================================================
# TOTAL
# ============================================================

dados_calculo = (
    dados_historico_salvar.copy()
)

dados_calculo["Total"] = (
    dados_calculo["Quantidade"]
    * dados_calculo["Valor"]
)

total_vendas = dados_calculo[
    "Total"
].sum()


# ============================================================
# LOG
# ============================================================

registrar_log(
    arquivos_processados=len(
        arquivos_validos
    ),
    arquivos_rejeitados=len(
        arquivos_rejeitados
    ),
    linhas_novas=linhas_novas,
    linhas_consolidadas=
        linhas_consolidadas,
    total_vendas=total_vendas
)


# ============================================================
# RESULTADO
# ============================================================

print()

print(
    "=" * 50
)

print(
    "AUTOMAÇÃO CONCLUÍDA!"
)

print(
    "=" * 50
)

print(
    f"Arquivos encontrados: "
    f"{len(arquivos)}"
)

print(
    f"Arquivos processados: "
    f"{len(arquivos_validos)}"
)

print(
    f"Arquivos rejeitados: "
    f"{len(arquivos_rejeitados)}"
)

print(
    f"Linhas lidas: "
    f"{linhas_lidas_total}"
)

print(
    f"Linhas novas: "
    f"{linhas_novas}"
)

print(
    f"Duplicidades ignoradas: "
    f"{duplicidades_ignoradas}"
)

print(
    f"Linhas consolidadas: "
    f"{linhas_consolidadas}"
)

print(
    f"Total das vendas: "
    f"R$ {total_vendas:.2f}"
)

print()

print(
    f"Relatório: "
    f"{arquivo_relatorio}"
)

print(
    f"Histórico: "
    f"{arquivo_historico}"
)

print(
    f"Log: "
    f"{arquivo_log}"
)

print(
    f"Controle de arquivos: "
    f"{arquivo_controle}"
)

print(
    f"Processados: "
    f"{pasta_processados}"
)

print()
